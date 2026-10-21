from __future__ import annotations

import sys

import docx
import pymupdf
import pytest

from docuforge_engine import errors
from docuforge_engine.convert import pdf_to_docx as module
from docuforge_engine.convert.pdf_to_docx import pdf_to_docx
from docuforge_engine.convert.pdf_to_docx_simple import _clean, build_simple_docx
from docuforge_engine.errors import EngineError
from docuforge_engine.safe_output import OutputTransaction, PendingOutput

from .conftest import leftover_temp_files


def texts(path):
    return [p.text for p in docx.Document(str(path)).paragraphs]


def write_html_pdf(path, html, pages=1):
    story = pymupdf.Story(html)
    writer = pymupdf.DocumentWriter(str(path))
    more = True
    while more:
        device = writer.begin_page(pymupdf.paper_rect("a4"))
        more, _ = story.place(pymupdf.paper_rect("a4") + (54, 54, -54, -54))
        story.draw(device)
        writer.end_page()
    writer.close()
    return path


# ---- the simple converter on its own ----


def test_the_simple_converter_keeps_text_formatting_bullets_and_pictures(tmp_path, report_pdf, make_ctx):
    output = tmp_path / "simple.docx"
    skipped = build_simple_docx(report_pdf, [1], output, make_ctx())

    document = docx.Document(str(output))
    text = "\n".join(p.text for p in document.paragraphs)
    for expected in ("Quarterly Report", "bold and italic paragraph", "First point", "Second point", "Closing paragraph"):
        assert expected in text
    runs = [run for p in document.paragraphs for run in p.runs]
    assert any(run.bold and "bold" in run.text for run in runs)
    assert any(run.italic and "italic" in run.text for run in runs)
    assert len(document.inline_shapes) == 1
    assert skipped == 0
    # bullets are separate paragraphs
    assert sum(1 for t in texts(output) if t.strip().endswith(("First point", "Second point"))) == 2


def test_each_pdf_page_ends_with_a_page_break(tmp_path, make_pdf, make_ctx):
    output = tmp_path / "pages.docx"
    build_simple_docx(make_pdf("doc", 3), [1, 2, 3], output, make_ctx())
    xml = docx.Document(str(output)).element.xml
    assert xml.count('w:type="page"') == 2
    assert all(f"doc-{n}" in "\n".join(texts(output)) for n in (1, 2, 3))


def test_right_to_left_paragraphs_are_marked(tmp_path, make_ctx):
    source = write_html_pdf(tmp_path / "urdu.pdf", "<p>English line</p><p dir='rtl'>اردو زبان میں لکھا ہوا متن</p>")
    output = tmp_path / "urdu.docx"
    build_simple_docx(source, [1], output, make_ctx())
    document = docx.Document(str(output))
    marked = [p for p in document.paragraphs if "<w:bidi" in p._p.xml]
    assert len(marked) == 1
    assert "English line" in document.paragraphs[0].text


def test_characters_xml_cannot_hold_are_removed():
    assert _clean("a\x00b\x01c\x0bd\ufffee") == "abcde"
    assert _clean("tab\tand\nnewline") == "tab\tand\nnewline"


# ---- falling back when the layout converter cannot be used ----


def run(source, output, make_ctx, **args):
    return pdf_to_docx({"input": str(source), "output": str(output), **args}, make_ctx())


def test_a_missing_layout_converter_falls_back_and_says_why(tmp_path, report_pdf, make_ctx, monkeypatch):
    monkeypatch.setitem(sys.modules, "pdf2docx", None)  # makes `from pdf2docx import ...` fail
    output = tmp_path / "out.docx"
    result = run(report_pdf, output, make_ctx)

    assert "simpler method" in result["note"] and "ModuleNotFoundError" in result["note"]
    assert "not tables or columns" in result["note"]
    assert "Quarterly Report" in "\n".join(texts(output))
    assert leftover_temp_files(tmp_path) == []


def test_a_layout_converter_that_crashes_falls_back(tmp_path, report_pdf, make_ctx, monkeypatch):
    import pdf2docx

    def explode(self, *args, **kwargs):
        raise KeyError("list index out of range")

    monkeypatch.setattr(pdf2docx.Converter, "convert", explode)
    output = tmp_path / "out.docx"
    result = run(report_pdf, output, make_ctx)
    assert "simpler method" in result["note"] and "KeyError" in result["note"]
    assert "Closing paragraph" in "\n".join(texts(output))


def test_a_layout_result_that_is_not_a_real_word_file_also_falls_back(tmp_path, report_pdf, make_ctx, monkeypatch):
    import pdf2docx

    def write_garbage(self, filename, *args, **kwargs):
        open(filename, "wb").write(b"PK\x03\x04 broken")

    monkeypatch.setattr(pdf2docx.Converter, "convert", write_garbage)
    output = tmp_path / "out.docx"
    result = run(report_pdf, output, make_ctx)
    assert "simpler method" in result["note"]
    assert "Quarterly Report" in "\n".join(texts(output))


def test_the_layout_converter_is_still_used_when_it_works(tmp_path, report_pdf, make_ctx):
    result = run(report_pdf, tmp_path / "out.docx", make_ctx)
    assert "layout was rebuilt" in result["note"]


def test_if_both_methods_fail_the_reason_is_in_the_message(tmp_path, report_pdf, make_ctx, monkeypatch):
    monkeypatch.setitem(sys.modules, "pdf2docx", None)

    def broken(*args, **kwargs):
        raise RuntimeError("python-docx exploded")

    monkeypatch.setattr(module, "build_simple_docx", broken)
    output = tmp_path / "out.docx"
    with pytest.raises(EngineError) as error:
        run(report_pdf, output, make_ctx)
    assert error.value.code == errors.CONVERSION_FAILED
    assert "RuntimeError: python-docx exploded" in error.value.message
    assert not output.exists()
    assert leftover_temp_files(tmp_path) == []


# ---- the safe-output check no longer lets a bad result escape as a crash ----


def test_a_result_that_is_not_a_real_word_file_is_rejected_cleanly(tmp_path, make_ctx):
    output = tmp_path / "out.docx"
    with pytest.raises(EngineError) as error:
        with OutputTransaction(make_ctx()) as transaction:
            temp = transaction.new_temp(output)
            temp.write_bytes(b"PK\x03\x04 this is not a real docx")
            transaction.commit([PendingOutput(temp, output, kind="docx")], overwrite=True)
    assert error.value.code == errors.OUTPUT_INVALID
    assert not output.exists()
