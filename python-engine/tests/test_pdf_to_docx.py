from __future__ import annotations

import io

import docx
import pymupdf
import pytest
from PIL import Image

from docuforge_engine import errors
from docuforge_engine.convert.pdf_to_docx import pdf_to_docx
from docuforge_engine.errors import EngineError

from .conftest import file_hash, leftover_temp_files


def convert(source, output, make_ctx, **args):
    return pdf_to_docx({"input": str(source), "output": str(output), **args}, make_ctx())


def test_text_formatting_bullets_and_pictures_survive(tmp_path, report_pdf, make_ctx):
    output = tmp_path / "report.docx"
    result = convert(report_pdf, output, make_ctx)

    document = docx.Document(str(output))
    text = "\n".join(p.text for p in document.paragraphs)
    for expected in ("Quarterly Report", "This is a bold and italic paragraph.", "First point", "Closing paragraph."):
        assert expected in text
    bold = [run.text.strip() for p in document.paragraphs for run in p.runs if run.bold]
    assert "bold" in bold
    assert len(document.inline_shapes) == 1
    assert "layout was rebuilt" in result["note"]
    assert result["pageCount"] == 1
    assert leftover_temp_files(tmp_path) == []


def test_a_page_range_limits_the_pages_converted(tmp_path, make_pdf, make_ctx):
    source = make_pdf("doc", 4)
    output = tmp_path / "doc.docx"
    result = convert(source, output, make_ctx, ranges=[[2, 3]])
    text = "\n".join(p.text for p in docx.Document(str(output)).paragraphs)
    assert "doc-2" in text and "doc-3" in text
    assert "doc-1" not in text and "doc-4" not in text
    assert result["pageCount"] == 2


def test_a_scanned_pdf_is_explained_and_nothing_is_saved(tmp_path, make_ctx):
    scan = tmp_path / "scan.pdf"
    with pymupdf.open() as document:
        page = document.new_page()
        pixmap = pymupdf.Pixmap(pymupdf.csRGB, pymupdf.IRect(0, 0, 50, 50), False)
        page.insert_image(pymupdf.Rect(0, 0, 200, 200), pixmap=pixmap)
        document.save(scan)
    output = tmp_path / "scan.docx"
    with pytest.raises(EngineError) as error:
        convert(scan, output, make_ctx)
    assert error.value.code == errors.NO_TEXT_FOUND
    assert "OCR" in error.value.message
    assert not output.exists()


def test_the_original_is_not_modified(tmp_path, report_pdf, make_ctx):
    before = file_hash(report_pdf)
    convert(report_pdf, tmp_path / "o.docx", make_ctx)
    assert file_hash(report_pdf) == before


@pytest.mark.parametrize(
    ("kind", "code"),
    [("encrypted", errors.PDF_ENCRYPTED), ("corrupt", errors.PDF_CORRUPT), ("missing", errors.INPUT_NOT_FOUND)],
)
def test_bad_inputs_fail_clearly(tmp_path, make_pdf, make_ctx, kind, code):
    if kind == "encrypted":
        source = make_pdf("locked", 1, password="pw")
    elif kind == "corrupt":
        source = tmp_path / "broken.pdf"
        source.write_bytes(b"nope")
    else:
        source = tmp_path / "missing.pdf"
    output = tmp_path / "o.docx"
    with pytest.raises(EngineError) as error:
        convert(source, output, make_ctx)
    assert error.value.code == code
    assert not output.exists()
    assert leftover_temp_files(tmp_path) == []


def test_invalid_ranges_and_cancel(tmp_path, make_pdf, make_ctx):
    source = make_pdf("doc", 2)
    with pytest.raises(EngineError) as error:
        convert(source, tmp_path / "o.docx", make_ctx, ranges=[[1, 9]])
    assert error.value.code == errors.INVALID_PAGE_RANGE

    output = tmp_path / "c.docx"
    with pytest.raises(EngineError) as error:
        pdf_to_docx({"input": str(source), "output": str(output)}, make_ctx(cancelled=True))
    assert error.value.code == errors.CANCELLED
    assert not output.exists()
    assert leftover_temp_files(tmp_path) == []
