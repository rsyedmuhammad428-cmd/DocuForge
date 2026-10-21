from __future__ import annotations

import zipfile

import pymupdf
import pytest

from docuforge_engine import errors
from docuforge_engine.convert import document_to_pdf as module
from docuforge_engine.convert.document_to_pdf import document_to_pdf, tools_status
from docuforge_engine.errors import EngineError
from docuforge_engine.tools import find_libreoffice

from .conftest import leftover_temp_files

needs_libreoffice = pytest.mark.skipif(find_libreoffice() is None, reason="LibreOffice is not installed")

_CONTENT_TYPES = (
    '<?xml version="1.0" encoding="UTF-8"?><Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
    '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
    '<Default Extension="xml" ContentType="application/xml"/>'
    '<Override PartName="/word/document.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/>'
    "</Types>"
)
_RELS = (
    '<?xml version="1.0" encoding="UTF-8"?><Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
    '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="word/document.xml"/>'
    "</Relationships>"
)


def make_docx(path, text):
    body = (
        '<?xml version="1.0" encoding="UTF-8"?><w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
        f"<w:body><w:p><w:r><w:t>{text}</w:t></w:r></w:p></w:body></w:document>"
    )
    with zipfile.ZipFile(path, "w") as archive:
        archive.writestr("[Content_Types].xml", _CONTENT_TYPES)
        archive.writestr("_rels/.rels", _RELS)
        archive.writestr("word/document.xml", body)
    return path


@needs_libreoffice
def test_a_word_document_becomes_a_pdf(tmp_path, make_ctx):
    source = make_docx(tmp_path / "letter.docx", "Hello DocuForge")
    output = tmp_path / "letter.pdf"
    result = document_to_pdf({"input": str(source), "output": str(output)}, make_ctx())

    assert result["pageCount"] == 1
    with pymupdf.open(output) as document:
        assert "Hello DocuForge" in document[0].get_text()
    assert leftover_temp_files(tmp_path) == []


@needs_libreoffice
def test_a_damaged_document_fails_clearly(tmp_path, make_ctx):
    source = tmp_path / "broken.docx"
    source.write_bytes(b"this is not a word document")
    output = tmp_path / "broken.pdf"
    with pytest.raises(EngineError) as error:
        document_to_pdf({"input": str(source), "output": str(output)}, make_ctx())
    assert error.value.code == errors.CONVERSION_FAILED
    assert not output.exists()


@needs_libreoffice
def test_cancel_stops_libreoffice_and_leaves_nothing(tmp_path, make_ctx):
    source = make_docx(tmp_path / "letter.docx", "Hello")
    output = tmp_path / "letter.pdf"
    with pytest.raises(EngineError) as error:
        document_to_pdf({"input": str(source), "output": str(output)}, make_ctx(cancelled=True))
    assert error.value.code == errors.CANCELLED
    assert not output.exists()
    assert leftover_temp_files(tmp_path) == []


def rich_docx(path):
    """A Word file made with python-docx: heading, formatting, a list, a table, a picture, Urdu."""
    import io

    import docx
    from docx.shared import Inches
    from PIL import Image

    document = docx.Document()
    document.add_heading("Meeting Notes", level=1)
    paragraph = document.add_paragraph("Plain text, ")
    paragraph.add_run("bold").bold = True
    document.add_paragraph("First item", style="List Bullet")
    table = document.add_table(rows=2, cols=2)
    table.style = "Table Grid"
    table.cell(0, 0).text, table.cell(0, 1).text = "Name", "Score"
    table.cell(1, 0).text, table.cell(1, 1).text = "Ali", "90"
    picture = io.BytesIO()
    Image.new("RGB", (300, 120), (200, 60, 60)).save(picture, "PNG")
    picture.seek(0)
    document.add_picture(picture, width=Inches(2))
    document.add_paragraph("اردو متن کی مثال")
    document.save(path)
    return path


def test_docx_converts_without_libreoffice(tmp_path, make_ctx, monkeypatch):
    monkeypatch.setattr(module, "find_libreoffice", lambda: None)
    source = rich_docx(tmp_path / "notes.docx")
    output = tmp_path / "notes.pdf"
    result = document_to_pdf({"input": str(source), "output": str(output)}, make_ctx())

    assert "built-in converter" in result["note"]
    with pymupdf.open(output) as document:
        text = document[0].get_text()
        assert "Meeting Notes" in text and "Plain text, bold" in text and "First item" in text
        assert "Name" in text and "Ali" in text and "90" in text
        assert any("\u0600" <= c <= "\u06ff" or "\ufb50" <= c <= "\ufeff" for c in text)  # Urdu
        assert len(document[0].get_images()) == 1
    assert leftover_temp_files(tmp_path) == []


def test_the_basic_method_can_be_chosen_even_when_libreoffice_exists(tmp_path, make_ctx):
    source = rich_docx(tmp_path / "notes.docx")
    result = document_to_pdf(
        {"input": str(source), "output": str(tmp_path / "o.pdf"), "method": "basic"}, make_ctx()
    )
    assert "built-in converter" in result["note"]


@needs_libreoffice
def test_libreoffice_is_preferred_for_docx_when_installed(tmp_path, make_ctx):
    source = make_docx(tmp_path / "letter.docx", "Hello DocuForge")
    result = document_to_pdf({"input": str(source), "output": str(tmp_path / "o.pdf")}, make_ctx())
    assert "LibreOffice" in result["note"]


def test_older_formats_need_libreoffice_and_say_how_to_get_around_it(tmp_path, make_ctx, monkeypatch):
    monkeypatch.setattr(module, "find_libreoffice", lambda: None)
    source = tmp_path / "old.rtf"
    source.write_bytes(b"{\\rtf1\\ansi Hello}")
    with pytest.raises(EngineError) as error:
        document_to_pdf({"input": str(source), "output": str(tmp_path / "o.pdf")}, make_ctx())
    assert error.value.code == errors.TOOL_NOT_FOUND
    assert "Save the file as .docx" in error.value.message
    assert tools_status({}, make_ctx())["libreoffice"] is False


def test_a_damaged_docx_fails_before_any_converter_runs(tmp_path, make_ctx, monkeypatch):
    monkeypatch.setattr(module, "find_libreoffice", lambda: None)
    source = tmp_path / "broken.docx"
    source.write_bytes(b"not a word document")
    with pytest.raises(EngineError) as error:
        document_to_pdf({"input": str(source), "output": str(tmp_path / "o.pdf")}, make_ctx())
    assert error.value.code == errors.CONVERSION_FAILED


def test_the_builtin_converter_flows_long_documents_onto_more_pages(tmp_path, make_ctx, monkeypatch):
    import docx

    monkeypatch.setattr(module, "find_libreoffice", lambda: None)
    document = docx.Document()
    for number in range(200):
        document.add_paragraph(f"Paragraph number {number}")
    document.save(tmp_path / "long.docx")
    result = document_to_pdf(
        {"input": str(tmp_path / "long.docx"), "output": str(tmp_path / "long.pdf")}, make_ctx()
    )
    assert result["pageCount"] >= 3


def test_the_builtin_converter_can_be_cancelled(tmp_path, make_ctx, monkeypatch):
    monkeypatch.setattr(module, "find_libreoffice", lambda: None)
    source = rich_docx(tmp_path / "notes.docx")
    output = tmp_path / "o.pdf"
    with pytest.raises(EngineError) as error:
        document_to_pdf({"input": str(source), "output": str(output)}, make_ctx(cancelled=True))
    assert error.value.code == errors.CANCELLED
    assert not output.exists()
    assert leftover_temp_files(tmp_path) == []


def test_unsupported_types_missing_files_and_same_file(tmp_path, make_ctx):
    text = tmp_path / "a.txt"
    text.write_text("x")
    with pytest.raises(EngineError) as error:
        document_to_pdf({"input": str(text), "output": str(tmp_path / "o.pdf")}, make_ctx())
    assert error.value.code == errors.CONVERSION_FAILED
    with pytest.raises(EngineError) as error:
        document_to_pdf({"input": str(tmp_path / "nope.docx"), "output": str(tmp_path / "o.pdf")}, make_ctx())
    assert error.value.code == errors.INPUT_NOT_FOUND
    with pytest.raises(EngineError) as error:
        document_to_pdf({"input": str(text), "output": str(text)}, make_ctx())
    assert error.value.code == errors.OUTPUT_IS_INPUT


def test_file_signatures_are_checked_before_conversion(tmp_path):
    from docuforge_engine.convert.document_to_pdf import looks_like_document

    rtf = tmp_path / "a.rtf"
    rtf.write_bytes(b"{\\rtf1\\ansi Hello}")
    doc = tmp_path / "old.doc"
    doc.write_bytes(bytes.fromhex("D0CF11E0A1B11AE1") + b"\x00" * 16)
    fake_docx = tmp_path / "fake.docx"
    fake_docx.write_bytes(b"plain text pretending to be a document")
    empty_zip = tmp_path / "empty.docx"
    with zipfile.ZipFile(empty_zip, "w") as archive:
        archive.writestr("readme.txt", "not a Word file")

    assert looks_like_document(rtf)
    assert looks_like_document(doc)
    assert looks_like_document(make_docx(tmp_path / "ok.docx", "x"))
    assert not looks_like_document(fake_docx)
    assert not looks_like_document(empty_zip)
    assert not looks_like_document(tmp_path / "missing.docx")
