"""Builds a small corpus of realistic Word documents and their PDFs (via LibreOffice).

Development tool, not shipped. Each document exercises something PDF to Word usually gets wrong.
Usage: python tools/fidelity/make_corpus.py OUTPUT_DIR
"""

from __future__ import annotations

import io
import subprocess
import sys
from pathlib import Path

import docx
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image

LOREM = (
    "The quick brown fox jumps over the lazy dog while the committee reviews the quarterly figures and "
    "considers how each region performed against the targets that were set at the start of the year. "
)


def picture(width=600, height=300, color=(40, 120, 200)):
    buffer = io.BytesIO()
    image = Image.new("RGB", (width, height), color)
    for x in range(0, width, 40):
        for y in range(0, height, 40):
            image.paste((255 - color[0], 255 - color[1], color[2]), (x, y, x + 20, y + 20))
    image.save(buffer, "PNG")
    buffer.seek(0)
    return buffer


def shade(cell, hex_color):
    properties = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:val"), "clear")
    shading.set(qn("w:fill"), hex_color)
    properties.append(shading)


def report(path):
    d = docx.Document()
    section = d.sections[0]
    section.header.paragraphs[0].text = "ACME Corporation - Quarterly Report"
    section.footer.paragraphs[0].text = "Confidential - do not distribute"
    d.add_heading("Quarterly Report", 0)
    d.add_heading("1. Summary", 1)
    p = d.add_paragraph(LOREM * 2)
    p.runs[0].font.size = Pt(11)
    d.add_paragraph("First finding", style="List Bullet")
    d.add_paragraph("Second finding", style="List Bullet")
    d.add_paragraph("Step one", style="List Number")
    d.add_paragraph("Step two", style="List Number")
    d.add_heading("2. Results", 1)
    table = d.add_table(rows=5, cols=4)
    table.style = "Table Grid"
    for r in range(5):
        for c in range(4):
            table.cell(r, c).text = f"Header {c}" if r == 0 else f"R{r}C{c}"
            if r == 0:
                shade(table.cell(r, c), "D9E2F3")
    d.add_paragraph()
    d.add_picture(picture(), width=Inches(4))
    d.add_paragraph("Figure 1: Results by region").alignment = WD_ALIGN_PARAGRAPH.CENTER
    d.add_heading("3. Details", 1)
    for _ in range(6):
        d.add_paragraph(LOREM * 3)
    d.save(path)


def formatted_letter(path):
    d = docx.Document()
    a = d.add_paragraph()
    a.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    a.add_run("Jane Smith\n12 Long Street\nLondon").font.size = Pt(10)
    d.add_paragraph()
    t = d.add_paragraph()
    run = t.add_run("Re: Application for the position")
    run.bold = True
    run.font.size = Pt(14)
    run.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)
    body = d.add_paragraph()
    body.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    body.paragraph_format.line_spacing = 1.5
    body.add_run(LOREM * 3)
    mixed = d.add_paragraph()
    mixed.add_run("Plain, ")
    mixed.add_run("bold, ").bold = True
    mixed.add_run("italic, ").italic = True
    mixed.add_run("underlined, ").underline = True
    colored = mixed.add_run("colored.")
    colored.font.color.rgb = RGBColor(0xC0, 0x00, 0x00)
    center = d.add_paragraph("Yours sincerely")
    center.alignment = WD_ALIGN_PARAGRAPH.CENTER
    d.add_paragraph(LOREM * 2)
    d.save(path)


def two_columns(path):
    d = docx.Document()
    d.add_heading("A Two Column Article", 0)
    section = d.sections[0]
    columns = OxmlElement("w:cols")
    columns.set(qn("w:num"), "2")
    columns.set(qn("w:space"), "720")
    section._sectPr.append(columns)
    for i in range(10):
        d.add_heading(f"Topic {i + 1}", 2)
        d.add_paragraph(LOREM * 2)
    d.save(path)


def table_heavy(path):
    d = docx.Document()
    d.add_heading("Price List", 1)
    table = d.add_table(rows=1, cols=4)
    table.style = "Light Grid Accent 1"
    for i, h in enumerate(["Item", "Qty", "Unit", "Total"]):
        table.rows[0].cells[i].text = h
    for r in range(15):
        row = table.add_row().cells
        row[0].text = f"Product {r}"
        row[1].text = str(r + 1)
        row[2].text = f"{(r + 1) * 1.25:.2f}"
        row[3].text = f"{(r + 1) * (r + 1) * 1.25:.2f}"
    d.add_paragraph()
    merged = d.add_table(rows=3, cols=3)
    merged.style = "Table Grid"
    merged.alignment = WD_TABLE_ALIGNMENT.CENTER
    top = merged.cell(0, 0).merge(merged.cell(0, 2))
    top.text = "Merged heading cell"
    for r in (1, 2):
        for c in range(3):
            merged.cell(r, c).text = f"Cell {r},{c}"
    d.save(path)


def urdu(path):
    d = docx.Document()
    d.add_heading("Urdu document", 1)
    for _ in range(4):
        p = d.add_paragraph("یہ اردو زبان میں لکھا ہوا ایک نمونہ متن ہے جو دائیں سے بائیں لکھا جاتا ہے۔")
        p._p.get_or_add_pPr().append(OxmlElement("w:bidi"))
    d.add_paragraph("English line after Urdu text.")
    d.save(path)


def long_text(path):
    d = docx.Document()
    d.add_heading("Long Document", 0)
    for i in range(14):
        d.add_heading(f"Chapter {i + 1}", 1)
        for _ in range(4):
            d.add_paragraph(LOREM * 4)
    d.save(path)


def images(path):
    d = docx.Document()
    d.add_heading("Image Gallery", 1)
    for i, color in enumerate([(200, 60, 60), (60, 160, 90), (70, 90, 200)]):
        d.add_picture(picture(500, 220, color), width=Inches(3.5))
        d.add_paragraph(f"Caption for image {i + 1}")
    d.save(path)


def resume(path):
    d = docx.Document()
    d.add_heading("JOHN DOE", 0).alignment = WD_ALIGN_PARAGRAPH.CENTER
    d.add_paragraph("Software Engineer | john@example.com | +44 7000 000000").alignment = WD_ALIGN_PARAGRAPH.CENTER
    layout = d.add_table(rows=1, cols=2)
    left, right = layout.rows[0].cells
    left.text = "EXPERIENCE\nSenior Engineer, ACME (2022-2026)\nBuilt things.\n\nEngineer, Globex (2019-2022)\nFixed things."
    right.text = "SKILLS\nPython\nTypeScript\nSQL\n\nEDUCATION\nBSc Software Engineering"
    d.save(path)


DOCUMENTS = {
    "report": report,
    "letter": formatted_letter,
    "columns": two_columns,
    "tables": table_heavy,
    "urdu": urdu,
    "long": long_text,
    "images": images,
    "resume": resume,
}


def main(output: Path) -> None:
    output.mkdir(parents=True, exist_ok=True)
    for name, build in DOCUMENTS.items():
        word = output / f"{name}.docx"
        build(word)
        subprocess.run(
            ["soffice", "--headless", "--convert-to", "pdf", "--outdir", str(output), str(word)],
            capture_output=True, timeout=180, check=True,
        )
        print("made", name)


if __name__ == "__main__":
    main(Path(sys.argv[1]))
