"""Runs every kind of operation against the PACKAGED engine, the way the app does.

A packaged program can fail in ways the normal one does not (missing data files, missing hidden
imports), so this must pass before an installer is released.
Usage: python tools/smoke_packaged.py PATH_TO_docuforge-engine[.exe]
"""

from __future__ import annotations

import io
import json
import subprocess
import sys
import tempfile
import zipfile
from pathlib import Path

import docx
import pymupdf
from PIL import Image

ENGINE = sys.argv[1]
WORK = Path(tempfile.mkdtemp(prefix="smoke-"))


def call(op: str, args: dict, job: str) -> dict:
    process = subprocess.Popen([ENGINE], stdin=subprocess.PIPE, stdout=subprocess.PIPE, stderr=subprocess.PIPE, text=True, encoding="utf-8")
    process.stdin.write(json.dumps({"id": job, "op": op, "args": args}) + "\n")
    process.stdin.flush()
    out = process.stdout.read()
    err = process.stderr.read()
    process.stdin.close()
    process.wait(timeout=300)
    messages = [json.loads(line) for line in out.splitlines() if line.strip().startswith("{")]
    final = messages[-1] if messages else {"type": "error", "code": "NO_OUTPUT", "message": err[-300:]}
    final["_stderr"] = err
    return final


results: list[tuple[str, bool, str]] = []


def check(name: str, ok: bool, detail: str = "") -> None:
    results.append((name, ok, detail))
    print(("PASS  " if ok else "FAIL  ") + name + (f"   {detail}" if detail and not ok else ""))


def make_pdf(path: Path, pages: int = 2) -> Path:
    with pymupdf.open() as d:
        for n in range(1, pages + 1):
            d.new_page().insert_text((72, 72), f"Hello page {n} of the smoke test document")
        d.save(path)
    return path


pdf_a, pdf_b = make_pdf(WORK / "a.pdf"), make_pdf(WORK / "b.pdf", 3)

r = call("inspect", {"paths": [str(pdf_a)]}, "s1")
check("inspect", r["type"] == "result" and r["data"]["files"][0]["pageCount"] == 2, str(r.get("message")))

r = call("merge", {"inputs": [str(pdf_a), str(pdf_b)], "output": str(WORK / "m.pdf")}, "s2")
check("merge", r["type"] == "result" and r["data"]["pageCount"] == 5, str(r.get("message")))

r = call("split", {"input": str(pdf_b), "outputDir": str(WORK), "mode": "each"}, "s3")
check("split", r["type"] == "result" and len(r["data"]["outputPaths"]) == 3, str(r.get("message")))

r = call("thumbnails", {"input": str(pdf_a), "pages": [0, 1]}, "s4")
check("thumbnails", r["type"] == "result" and len(r["data"]["thumbnails"]) == 2, str(r.get("message")))

r = call("compress", {"input": str(WORK / "m.pdf"), "output": str(WORK / "m_small.pdf"), "preset": "light"}, "s5")
check("compress", r["type"] in ("result", "error") and (r["type"] == "result" or r["code"] == "NO_SIZE_GAIN"), str(r.get("message")))

r = call("pdf_to_images", {"input": str(pdf_a), "outputDir": str(WORK), "format": "png", "dpi": 100}, "s6")
check("pdf_to_images", r["type"] == "result", str(r.get("message")))

png = WORK / "i.png"
Image.new("RGB", (200, 100), (200, 50, 50)).save(png)
r = call("images_to_pdf", {"inputs": [str(png)], "output": str(WORK / "img.pdf")}, "s7")
check("images_to_pdf (Pillow)", r["type"] == "result", str(r.get("message")))

r = call("pdf_to_text", {"input": str(pdf_a), "output": str(WORK / "t.txt")}, "s8")
check("pdf_to_text", r["type"] == "result", str(r.get("message")))

r = call("text_to_pdf", {"input": str(WORK / "t.txt"), "output": str(WORK / "t.pdf")}, "s9")
check("text_to_pdf", r["type"] == "result", str(r.get("message")))

word = WORK / "w.docx"
d = docx.Document(); d.add_heading("Smoke", 1); d.add_paragraph("Word to PDF with the built-in converter."); d.save(word)
r = call("document_to_pdf", {"input": str(word), "output": str(WORK / "w.pdf"), "method": "basic"}, "s10")
check("document_to_pdf (mammoth, built-in)", r["type"] == "result" and "built-in" in r["data"].get("note", ""), str(r.get("message")))

r = call("pdf_to_docx", {"input": str(WORK / "m.pdf"), "output": str(WORK / "m.docx")}, "s11")
note = r.get("data", {}).get("note", "")
check("pdf_to_docx (layout converter: pdf2docx + OpenCV)", r["type"] == "result" and "layout was rebuilt" in note, f"{r.get('message')} {note[:160]} {r['_stderr'][-300:]}")
if r["type"] == "result":
    check("pdf_to_docx result opens in python-docx", len(docx.Document(str(WORK / "m.docx")).paragraphs) > 0)

r = call("tools_status", {}, "s12")
check("tools_status", r["type"] == "result" and "tesseract" in r["data"], str(r.get("message")))

r = call("explode", {}, "s13")
check("unknown operation is reported cleanly", r["type"] == "error" and r["code"] == "UNKNOWN_OPERATION")

# OCR needs Tesseract, which is an optional separate program
status = call("tools_status", {}, "s14")
tess = status.get("data", {}).get("tesseract", {})
if tess.get("found") and "eng" in tess.get("languages", []):
    with pymupdf.open() as src:
        page = src.new_page()
        page.insert_textbox(pymupdf.Rect(72, 72, 520, 300), "Invoice 2026 total due 1250", fontsize=18)
        pix = page.get_pixmap(dpi=200)
    scan = WORK / "scan.pdf"
    with pymupdf.open() as s:
        p = s.new_page(width=595, height=842); p.insert_image(p.rect, pixmap=pix); s.save(scan)
    r = call("ocr", {"input": str(scan), "output": str(WORK / "scan.txt"), "format": "text", "languages": ["eng"], "dpi": 300, "useExistingText": True, "pageMarkers": False}, "s15")
    check("ocr (external Tesseract)", r["type"] == "result" and "Invoice" in r["data"].get("preview", ""), str(r.get("message")))
else:
    print("SKIP  ocr (Tesseract not installed here)")

failed = [name for name, ok, _ in results if not ok]
print(f"\n{len(results) - len(failed)} of {len(results)} checks passed")
sys.exit(1 if failed else 0)
