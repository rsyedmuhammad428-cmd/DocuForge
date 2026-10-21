"""A simple PDF to Word converter that needs only PyMuPDF and python-docx.

It is the fallback for `pdf_to_docx` when the layout converter (pdf2docx, which needs OpenCV and
NumPy) cannot be used on this computer or fails on a particular file.

It keeps: text with bold, italic, size and color; paragraphs and bullet lists; centered lines;
right-to-left text (Urdu, Arabic); pictures; and page breaks.
It does NOT detect tables or columns: each block of text becomes its own paragraph, in reading order.
"""

from __future__ import annotations

import io
import re
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.shared import Emu, Pt, RGBColor

from ..context import JobContext
from ..pdf_tools.common import open_pdf

_BULLET_START = re.compile(r"^\s*([•·▪◦‣○●■□*\-–—]|\d{1,3}[.)])\s+")
_RTL_CHARACTERS = re.compile(r"[\u0590-\u08ff\ufb1d-\ufdff\ufe70-\ufeff]")
_XML_INVALID = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f\ufffe\uffff]")
_EMU_PER_POINT = 12700
_MARGIN_POINTS = 40


def _clean(text: str) -> str:
    """python-docx refuses characters that XML cannot hold; PDFs sometimes contain them."""
    return _XML_INVALID.sub("", text)


def _make_rtl(paragraph: Any) -> None:
    properties = paragraph._p.get_or_add_pPr()  # noqa: SLF001 - python-docx has no public API for this
    properties.append(OxmlElement("w:bidi"))
    for run in paragraph.runs:
        run_properties = run._r.get_or_add_rPr()  # noqa: SLF001
        run_properties.append(OxmlElement("w:rtl"))


def _split_paragraphs(block: dict[str, Any]) -> list[list[dict[str, Any]]]:
    """Joins a block's lines into paragraphs; a bullet or number starts a new one."""
    paragraphs: list[list[dict[str, Any]]] = [[]]
    for line in block.get("lines", []):
        spans = [span for span in line.get("spans", []) if span.get("text")]
        if not spans:
            continue
        line_text = "".join(span["text"] for span in spans)
        if paragraphs[-1] and _BULLET_START.match(line_text):
            paragraphs.append([])
        if paragraphs[-1]:
            paragraphs[-1].append({"text": " ", "flags": 0, "size": spans[0]["size"], "font": "", "color": 0})
        paragraphs[-1].extend(spans)
    return [paragraph for paragraph in paragraphs if paragraph]


def _add_text_block(document: Any, block: dict[str, Any], page_width: float) -> None:
    for spans in _split_paragraphs(block):
        text = _clean("".join(span["text"] for span in spans))
        if not text.strip():
            continue
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(4)

        for span in spans:
            part = _clean(span["text"])
            if not part:
                continue
            font_name = str(span.get("font", "")).lower()
            run = paragraph.add_run(part)
            run.bold = bool(span.get("flags", 0) & 16) or "bold" in font_name
            run.italic = bool(span.get("flags", 0) & 2) or "italic" in font_name or "oblique" in font_name
            run.font.size = Pt(max(6.0, min(72.0, round(float(span.get("size", 11)) * 2) / 2)))
            color = int(span.get("color", 0))
            if color:
                run.font.color.rgb = RGBColor((color >> 16) & 255, (color >> 8) & 255, color & 255)

        x0, _, x1, _ = block["bbox"]
        if abs((x0 + x1) / 2 - page_width / 2) < page_width * 0.03 and (x1 - x0) < page_width * 0.8:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        letters = [c for c in text if c.isalpha()]
        if letters and sum(1 for c in letters if _RTL_CHARACTERS.match(c)) > len(letters) / 2:
            _make_rtl(paragraph)


def build_simple_docx(
    source_path: Path, page_numbers: list[int], output_path: Path, ctx: JobContext
) -> int:
    """Writes the .docx. Returns how many pictures could not be copied."""
    skipped_pictures = 0
    document = Document()

    with open_pdf(source_path) as pdf:
        first = pdf.load_page(page_numbers[0] - 1).rect
        section = document.sections[0]
        section.page_width = Emu(int(first.width * _EMU_PER_POINT))
        section.page_height = Emu(int(first.height * _EMU_PER_POINT))
        for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
            setattr(section, side, Emu(_MARGIN_POINTS * _EMU_PER_POINT))

        for position, number in enumerate(page_numbers):
            ctx.check_cancelled()
            page = pdf.load_page(number - 1)
            usable_width = max(50.0, page.rect.width - 2 * _MARGIN_POINTS)
            for block in page.get_text("dict", sort=True).get("blocks", []):
                if block.get("type") == 0:
                    _add_text_block(document, block, page.rect.width)
                elif block.get("type") == 1:
                    x0, _, x1, _ = block["bbox"]
                    try:
                        document.add_picture(io.BytesIO(block["image"]), width=Pt(max(10.0, min(x1 - x0, usable_width))))
                    except Exception:  # noqa: BLE001 - python-docx cannot read every image type
                        skipped_pictures += 1
            if position < len(page_numbers) - 1:
                document.add_page_break()
            ctx.progress(2, 4)

    ctx.check_cancelled()
    document.save(str(output_path))
    return skipped_pictures


_EMU_PER_POINT = 12700
